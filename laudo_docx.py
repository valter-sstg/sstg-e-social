"""
Conversor da "story" do ReportLab (lista de flowables usada nos laudos PDF) para .docx.

Os laudos DRPS (gerar_laudo.py) e DRE/AEP (gerar_laudo_aep.py) montam o conteúdo uma
única vez como lista de flowables (Paragraph, Table, Spacer, Image, PageBreak...). Este
módulo percorre essa mesma lista e reproduz o conteúdo em Word com python-docx, de modo
que qualquer ajuste de texto nos build_* sai igual no PDF e no Word.
"""
import io
import os
import re
from html.parser import HTMLParser

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, Emu, RGBColor

from reportlab.lib import colors as rl_colors
from reportlab.platypus import (
    Paragraph, Table, Spacer, PageBreak, Image, KeepTogether, HRFlowable, CondPageBreak
)

FONTE = "Arial"
_PT = 12700  # EMU por ponto tipográfico

_ALINHAMENTO = {
    0: WD_ALIGN_PARAGRAPH.LEFT, 1: WD_ALIGN_PARAGRAPH.CENTER,
    2: WD_ALIGN_PARAGRAPH.RIGHT, 4: WD_ALIGN_PARAGRAPH.JUSTIFY,
    "LEFT": WD_ALIGN_PARAGRAPH.LEFT, "CENTER": WD_ALIGN_PARAGRAPH.CENTER,
    "CENTRE": WD_ALIGN_PARAGRAPH.CENTER, "RIGHT": WD_ALIGN_PARAGRAPH.RIGHT,
    "DECIMAL": WD_ALIGN_PARAGRAPH.RIGHT, "JUSTIFY": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
_VALIGN = {
    "TOP": WD_CELL_VERTICAL_ALIGNMENT.TOP,
    "MIDDLE": WD_CELL_VERTICAL_ALIGNMENT.CENTER,
    "BOTTOM": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM,
}


# ===================== UTILITÁRIOS =====================

def _hex(cor):
    """Cor ReportLab → 'RRGGBB' (None se transparente/ausente)."""
    if cor is None:
        return None
    if isinstance(cor, str):
        try:
            cor = rl_colors.toColor(cor)
        except Exception:
            return None
    if getattr(cor, "alpha", 1) == 0:
        return None
    return "%02X%02X%02X" % (round(cor.red * 255), round(cor.green * 255), round(cor.blue * 255))


def _rgb(cor):
    h = _hex(cor)
    return RGBColor.from_string(h) if h else None


def _sombrear(elemento_pr, hex_cor):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_cor)
    elemento_pr.append(shd)


def _margens_celula(celula, top, bottom, left, right):
    tc_pr = celula._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for lado, valor in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:w"), str(int(max(valor, 0) * 20)))  # pontos → twips
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


def _borda_celula(celula, lado, espessura_pt, hex_cor):
    tc_pr = celula._tc.get_or_add_tcPr()
    bordas = tc_pr.find(qn("w:tcBorders"))
    if bordas is None:
        bordas = OxmlElement("w:tcBorders")
        tc_pr.append(bordas)
    el = bordas.find(qn(f"w:{lado}"))
    if el is None:
        el = OxmlElement(f"w:{lado}")
        bordas.append(el)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(max(2, int(espessura_pt * 8))))  # oitavos de ponto
    el.set(qn("w:color"), hex_cor or "000000")


def _sem_bordas(tabela):
    tbl_pr = tabela._tbl.tblPr
    bordas = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "nil")
        bordas.append(el)
    tbl_pr.append(bordas)


def _fonte_run(run, nome_fonte, tamanho, cor, negrito=False, italico=False):
    run.font.name = FONTE
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONTE)
    run.font.size = Pt(tamanho)
    nome = (nome_fonte or "").lower()
    run.bold = negrito or "bold" in nome
    run.italic = italico or "oblique" in nome or "italic" in nome
    rgb = _rgb(cor)
    if rgb is not None:
        run.font.color.rgb = rgb


# ===================== PARÁGRAFOS (mini-markup do ReportLab) =====================

class _MarkupParser(HTMLParser):
    """Converte o mini-HTML do ReportLab (<b>, <i>, <u>, <br/>, <font>) em trechos formatados."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.trechos = []  # (texto, {b, i, u, cor, tamanho}) ou ("\n", None)
        self._pilha = [{}]

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if tag == "br":
            self.trechos.append(("\n", None))
            return
        fmt = dict(self._pilha[-1])
        if tag in ("b", "strong"):
            fmt["b"] = True
        elif tag in ("i", "em"):
            fmt["i"] = True
        elif tag == "u":
            fmt["u"] = True
        elif tag in ("font", "span"):
            a = dict(attrs)
            if a.get("color"):
                fmt["cor"] = a["color"]
            if a.get("size"):
                try:
                    fmt["tamanho"] = float(a["size"])
                except ValueError:
                    pass
        self._pilha.append(fmt)

    def handle_startendtag(self, tag, attrs):
        if tag.lower() == "br":
            self.trechos.append(("\n", None))

    def handle_endtag(self, tag):
        if tag.lower() != "br" and len(self._pilha) > 1:
            self._pilha.pop()

    def handle_data(self, data):
        data = re.sub(r"\s+", " ", data)
        if data:
            self.trechos.append((data, dict(self._pilha[-1])))


def _preencher_paragrafo(par, markup, estilo, alinhamento=None, cor_padrao=None,
                         fonte_padrao=None, tamanho_padrao=None):
    fonte = getattr(estilo, "fontName", None) or fonte_padrao or "Helvetica"
    tamanho = getattr(estilo, "fontSize", None) or tamanho_padrao or 9
    cor = getattr(estilo, "textColor", None) or cor_padrao

    parser = _MarkupParser()
    parser.feed(str(markup))
    parser.close()
    # remove espaço em branco nas bordas (como o ReportLab faz)
    trechos = parser.trechos
    if trechos and trechos[0][1] is not None:
        trechos[0] = (trechos[0][0].lstrip(), trechos[0][1])
    if trechos and trechos[-1][1] is not None:
        trechos[-1] = (trechos[-1][0].rstrip(), trechos[-1][1])

    for texto, fmt in trechos:
        if fmt is None:
            par.add_run().add_break()
            continue
        if not texto:
            continue
        run = par.add_run(texto)
        _fonte_run(run, fonte, fmt.get("tamanho", tamanho), fmt.get("cor", cor),
                   fmt.get("b", False), fmt.get("i", False))
        if fmt.get("u"):
            run.underline = True

    pf = par.paragraph_format
    al = alinhamento if alinhamento is not None else getattr(estilo, "alignment", 0)
    pf.alignment = _ALINHAMENTO.get(al, WD_ALIGN_PARAGRAPH.LEFT)
    pf.space_before = Pt(getattr(estilo, "spaceBefore", 0) or 0)
    pf.space_after = Pt(getattr(estilo, "spaceAfter", 0) or 0)
    leading = getattr(estilo, "leading", None)
    if leading:
        pf.line_spacing = Pt(leading)
    recuo = getattr(estilo, "leftIndent", 0) or 0
    if recuo:
        pf.left_indent = Pt(recuo)
    return par


# ===================== CONVERSOR =====================

class _Conversor:
    # ---- dispatcher -------------------------------------------------------
    def adicionar(self, container, flowable, largura_emu):
        if isinstance(flowable, (list, tuple)):
            for f in flowable:
                self.adicionar(container, f, largura_emu)
        elif isinstance(flowable, KeepTogether):
            raiz = container._tc if hasattr(container, "_tc") else container.element.body
            antes = len(raiz)
            self.adicionar(container, getattr(flowable, "_content", []), largura_emu)
            # equivalente Word do KeepTogether: "manter com o próximo" em tudo que foi gerado
            # (evita título de seção sozinho no pé da página)
            for el in list(raiz)[antes:]:
                for p in ([el] if el.tag == qn("w:p") else el.iter(qn("w:p"))):
                    p_pr = p.get_or_add_pPr()
                    if p_pr.find(qn("w:keepNext")) is None:
                        p_pr.append(OxmlElement("w:keepNext"))
        elif isinstance(flowable, Paragraph):
            par = container.add_paragraph()
            _preencher_paragrafo(par, flowable.text, flowable.style)
        elif isinstance(flowable, Table):
            self._tabela(container, flowable, largura_emu)
        elif isinstance(flowable, Image):
            self._imagem(container.add_paragraph(), flowable, largura_emu)
        elif isinstance(flowable, PageBreak):
            container.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif isinstance(flowable, CondPageBreak):
            pass  # quebra condicional não tem equivalente direto; o Word pagina sozinho
        elif isinstance(flowable, HRFlowable):
            self._linha_horizontal(container.add_paragraph(), flowable)
        elif isinstance(flowable, Spacer):
            par = container.add_paragraph()
            pf = par.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.line_spacing = Pt(max(float(getattr(flowable, "height", 0) or 0), 1))
        elif isinstance(flowable, str):
            if flowable.strip():
                par = container.add_paragraph()
                _preencher_paragrafo(par, flowable, None)

    # ---- imagem -----------------------------------------------------------
    def _imagem(self, par, img, largura_max_emu):
        origem = img.filename
        if not (isinstance(origem, str) and os.path.exists(origem)):
            # imagem em memória (ex.: gráfico matplotlib em BytesIO): o ReportLab guarda
            # um ImageReader com a imagem PIL — exporta para PNG
            leitor = getattr(img, "_img", None)
            pil = getattr(leitor, "_image", None)
            if pil is None:
                return
            origem = io.BytesIO()
            pil.save(origem, format="PNG")
            origem.seek(0)
        largura = Emu(int(img.drawWidth * _PT))
        altura = Emu(int(img.drawHeight * _PT))
        if largura_max_emu and largura > largura_max_emu:
            fator = largura_max_emu / largura
            largura, altura = Emu(int(largura * fator)), Emu(int(altura * fator))
        par.add_run().add_picture(origem, width=largura, height=altura)
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _linha_horizontal(self, par, hr):
        p_pr = par._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), str(max(2, int((getattr(hr, "lineWidth", 1) or 1) * 8))))
        bottom.set(qn("w:color"), _hex(getattr(hr, "color", None)) or "CCCCCC")
        bdr.append(bottom)
        p_pr.append(bdr)

    # ---- tabela -----------------------------------------------------------
    def _tabela(self, container, t, largura_disp_emu):
        linhas = t._cellvalues
        n_lin = len(linhas)
        n_col = max((len(l) for l in linhas), default=0)
        if not n_lin or not n_col:
            return

        # larguras (pontos → EMU); None = divide o espaço restante
        larguras = list(getattr(t, "_argW", None) or t._colWidths or [None] * n_col)
        larguras = [w if isinstance(w, (int, float)) else None for w in larguras][:n_col]
        larguras += [None] * (n_col - len(larguras))
        fixas = sum(w * _PT for w in larguras if w)
        livres = larguras.count(None)
        resto = max(largura_disp_emu - fixas, Cm(1) * livres) if livres else 0
        larg_emu = [int(w * _PT) if w else int(resto / livres) for w in larguras]
        total = sum(larg_emu)
        if total > largura_disp_emu:  # não ultrapassa a margem do Word
            larg_emu = [int(w * largura_disp_emu / total) for w in larg_emu]

        tabela = container.add_table(rows=n_lin, cols=n_col)
        tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
        tabela.autofit = False
        _sem_bordas(tabela)
        for r in range(n_lin):
            for c in range(n_col):
                tabela.cell(r, c).width = Emu(larg_emu[c])

        def _norm(sc, sr, ec, er):
            sc, ec = (sc + n_col if sc < 0 else sc), (ec + n_col if ec < 0 else ec)
            sr, er = (sr + n_lin if sr < 0 else sr), (er + n_lin if er < 0 else er)
            return sc, sr, min(ec, n_col - 1), min(er, n_lin - 1)

        # fundos
        fundo = [[None] * n_col for _ in range(n_lin)]
        for cmd in getattr(t, "_bkgrndcmds", []):
            op, (sc, sr), (ec, er), arg = cmd[0], cmd[1], cmd[2], cmd[3]
            sc, sr, ec, er = _norm(sc, sr, ec, er)
            for r in range(sr, er + 1):
                for c in range(sc, ec + 1):
                    if op == "BACKGROUND":
                        fundo[r][c] = arg
                    elif op == "ROWBACKGROUNDS" and arg:
                        fundo[r][c] = arg[(r - sr) % len(arg)]
                    elif op == "COLBACKGROUNDS" and arg:
                        fundo[r][c] = arg[(c - sc) % len(arg)]

        # bordas
        for cmd in getattr(t, "_linecmds", []):
            op, (sc, sr), (ec, er), esp, cor = cmd[0], cmd[1], cmd[2], cmd[3], cmd[4]
            sc, sr, ec, er = _norm(sc, sr, ec, er)
            hx = _hex(cor)
            if not hx or not esp:
                continue
            for r in range(sr, er + 1):
                for c in range(sc, ec + 1):
                    cel = tabela.cell(r, c)
                    lados = []
                    if op == "GRID":
                        lados = ["top", "bottom", "start", "end"]
                    elif op in ("BOX", "OUTLINE"):
                        lados = [lado for lado, cond in (("top", r == sr), ("bottom", r == er),
                                                         ("start", c == sc), ("end", c == ec)) if cond]
                    elif op == "INNERGRID":
                        lados = [lado for lado, cond in (("top", r > sr), ("start", c > sc)) if cond]
                    elif op == "LINEBELOW":
                        lados = ["bottom"]
                    elif op == "LINEABOVE":
                        lados = ["top"]
                    elif op == "LINEBEFORE":
                        lados = ["start"]
                    elif op == "LINEAFTER":
                        lados = ["end"]
                    for lado in lados:
                        _borda_celula(cel, lado, esp, hx)

        # células (conteúdo + estilo)
        for r in range(n_lin):
            for c in range(n_col):
                valor = linhas[r][c] if c < len(linhas[r]) else ""
                cel = tabela.cell(r, c)
                estilo_cel = t._cellStyles[r][c]
                if fundo[r][c] is not None and _hex(fundo[r][c]):
                    _sombrear(cel._tc.get_or_add_tcPr(), _hex(fundo[r][c]))
                _margens_celula(cel, estilo_cel.topPadding, estilo_cel.bottomPadding,
                                estilo_cel.leftPadding, estilo_cel.rightPadding)
                cel.vertical_alignment = _VALIGN.get(str(estilo_cel.valign).upper(),
                                                     WD_CELL_VERTICAL_ALIGNMENT.TOP)
                larg_interna = max(larg_emu[c] - int((estilo_cel.leftPadding + estilo_cel.rightPadding) * _PT), Cm(1))
                self._conteudo_celula(cel, valor, estilo_cel, larg_interna)

        # mesclagens (SPAN) — feitas depois de preencher, mantendo o conteúdo da 1ª célula
        for cmd in getattr(t, "_spanCmds", []):
            (sc, sr), (ec, er) = cmd[1], cmd[2]
            sc, sr, ec, er = _norm(sc, sr, ec, er)
            if (sc, sr) == (ec, er):
                continue
            a = tabela.cell(sr, sc)
            for r in range(sr, er + 1):
                for c in range(sc, ec + 1):
                    if (r, c) != (sr, sc):
                        for p in list(tabela.cell(r, c)._tc.iter(qn("w:p"))):
                            p.getparent().remove(p)
                        tabela.cell(r, c)._tc.append(OxmlElement("w:p"))
            a.merge(tabela.cell(er, ec))
            # após o merge o python-docx concatena parágrafos vazios; remove os que sobraram
            ps = a.paragraphs
            for p in ps[1:]:
                if not p.text and not p._p.xpath(".//w:drawing") and len(ps) > 1:
                    p._p.getparent().remove(p._p)

        # cabeçalho repetido em cada página
        repetir = getattr(t, "repeatRows", 0)
        for r in range(min(repetir, n_lin) if isinstance(repetir, int) else 0):
            tr_pr = tabela.rows[r]._tr.get_or_add_trPr()
            el = OxmlElement("w:tblHeader")
            el.set(qn("w:val"), "true")
            tr_pr.append(el)

        # evita que o Word cole a próxima tabela nesta (parágrafo separador mínimo)
        sep = container.add_paragraph()
        sep.paragraph_format.space_after = Pt(0)
        sep.paragraph_format.space_before = Pt(0)
        sep.paragraph_format.line_spacing = Pt(1)

    def _conteudo_celula(self, cel, valor, estilo_cel, largura_emu):
        par_inicial = cel.paragraphs[0]
        usado = False
        itens = valor if isinstance(valor, (list, tuple)) else [valor]
        for item in itens:
            if item is None or (isinstance(item, str) and not item.strip()):
                continue
            if isinstance(item, Paragraph):
                par = par_inicial if not usado else cel.add_paragraph()
                _preencher_paragrafo(par, item.text, item.style)
            elif isinstance(item, str):
                par = par_inicial if not usado else cel.add_paragraph()
                _preencher_paragrafo(par, item, None,
                                     alinhamento=str(estilo_cel.alignment).upper(),
                                     cor_padrao=estilo_cel.color,
                                     fonte_padrao=estilo_cel.fontname,
                                     tamanho_padrao=estilo_cel.fontsize)
            elif isinstance(item, Image):
                par = par_inicial if not usado else cel.add_paragraph()
                self._imagem(par, item, largura_emu)
            elif isinstance(item, (Table, KeepTogether, list, tuple, Spacer, HRFlowable)):
                if isinstance(item, Spacer) and not usado:
                    continue
                self.adicionar(cel, item, largura_emu)
            else:
                continue
            usado = True
        # remove o parágrafo inicial se ficou vazio e há outro conteúdo
        if not usado:
            par_inicial.paragraph_format.space_after = Pt(0)
            par_inicial.paragraph_format.line_spacing = Pt(1)
        elif not par_inicial.text and not par_inicial._p.xpath(".//w:drawing") and len(cel._tc.xpath("./w:p|./w:tbl")) > 1:
            par_inicial._p.getparent().remove(par_inicial._p)
        # a última coisa numa célula do Word deve ser um parágrafo
        if cel._tc[-1].tag == qn("w:tbl"):
            p = cel.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(1)


# ===================== CABEÇALHO / RODAPÉ =====================

def _campo_pagina(run):
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = texto
        run._r.append(el)


def _cabecalho_rodape(secao, titulo_cabecalho, texto_rodape, cor_barra, cor_linha, logo_path):
    secao.different_first_page_header_footer = True  # capa sem cabeçalho, como no PDF
    largura = secao.page_width - secao.left_margin - secao.right_margin

    # cabeçalho: barra azul com (logo) título à esquerda e "Pág. N" à direita
    cab = secao.header
    cab.is_linked_to_previous = False
    t = cab.add_table(rows=1, cols=3 if logo_path else 2, width=largura)
    t.autofit = False
    _sem_bordas(t)
    celulas = t.rows[0].cells
    idx = 0
    if logo_path:
        c_logo = celulas[0]
        c_logo.width = Cm(2.8)
        p = c_logo.paragraphs[0]
        try:
            p.add_run().add_picture(logo_path, height=Cm(0.9))
        except Exception:
            pass
        idx = 1
    c_tit, c_pag = celulas[idx], celulas[idx + 1]
    c_pag.width = Cm(2.5)
    c_tit.width = largura - Cm(2.5) - (Cm(2.8) if logo_path else 0)
    r = c_tit.paragraphs[0].add_run(titulo_cabecalho)
    _fonte_run(r, "Helvetica-Bold", 8, rl_colors.white)
    p_pag = c_pag.paragraphs[0]
    p_pag.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _fonte_run(p_pag.add_run("Pág. "), "Helvetica", 8, rl_colors.white)
    r_num = p_pag.add_run()
    _fonte_run(r_num, "Helvetica", 8, rl_colors.white)
    _campo_pagina(r_num)
    for cel in celulas:
        _sombrear(cel._tc.get_or_add_tcPr(), _hex(cor_barra))
        _borda_celula(cel, "bottom", 3, _hex(cor_linha))
        cel.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _margens_celula(cel, 4, 4, 6, 6)
    cab.paragraphs[0]._p.getparent().remove(cab.paragraphs[0]._p)
    cab.add_paragraph().paragraph_format.line_spacing = Pt(1)

    # rodapé: barra azul com texto centralizado
    rod = secao.footer
    rod.is_linked_to_previous = False
    tr = rod.add_table(rows=1, cols=1, width=largura)
    _sem_bordas(tr)
    c = tr.rows[0].cells[0]
    _sombrear(c._tc.get_or_add_tcPr(), _hex(cor_barra))
    _margens_celula(c, 3, 3, 6, 6)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _fonte_run(p.add_run(texto_rodape), "Helvetica", 6.5, rl_colors.white)
    rod.paragraphs[0]._p.getparent().remove(rod.paragraphs[0]._p)
    rod.add_paragraph().paragraph_format.line_spacing = Pt(1)


# ===================== API =====================

def story_para_docx(story, *, titulo, autor, assunto, titulo_cabecalho, texto_rodape,
                    cor_barra, cor_linha, logo_path=None) -> bytes:
    """Converte a story do ReportLab num .docx (A4, mesmas margens do PDF) e retorna os bytes."""
    doc = Document()
    secao = doc.sections[0]
    secao.page_width, secao.page_height = Cm(21), Cm(29.7)
    secao.left_margin = secao.right_margin = Cm(1.5)
    secao.top_margin = Cm(2.2)
    secao.bottom_margin = Cm(1.5)
    secao.header_distance = Cm(0.5)
    secao.footer_distance = Cm(0.4)

    normal = doc.styles["Normal"]
    normal.font.name = FONTE
    normal.font.size = Pt(9)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONTE)
    normal.paragraph_format.space_after = Pt(0)

    # não esticar linhas terminadas em quebra manual (<br/>) em parágrafos justificados
    compat = doc.settings.element.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        doc.settings.element.append(compat)
    compat.insert(0, OxmlElement("w:doNotExpandShiftReturn"))

    doc.core_properties.title = titulo
    doc.core_properties.author = autor
    doc.core_properties.subject = assunto

    _cabecalho_rodape(secao, titulo_cabecalho, texto_rodape, cor_barra, cor_linha, logo_path)

    largura_util = secao.page_width - secao.left_margin - secao.right_margin
    _Conversor().adicionar(doc, story, largura_util)

    # remove o parágrafo vazio inicial criado pelo template padrão
    primeiro = doc.paragraphs[0] if doc.paragraphs else None
    if primeiro is not None and not primeiro.text and primeiro._p is doc.element.body[0]:
        primeiro._p.getparent().remove(primeiro._p)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
